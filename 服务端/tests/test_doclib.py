# -*- coding: utf-8 -*-
"""只读文件库：只读纪律、解析六态、变化感知。

三组测试对应这个模块仅有的三条纪律，**少钉一条这个模块就失去意义**：

1. 只读。扫描前后把目录完整比对一遍（文件名、大小、修改时间、内容摘要）。
   这条不靠人自觉，靠测试——「不小心写了个临时文件进客户资料目录」
   这种事在 code review 里基本看不出来。

2. 六态不许合并。尤其「需OCR」与「未解析」：前者是 PDF 没文字层、本机能办；
   后者是格式本身取不到、要换路子。合成一个状态，没人知道下一步该干什么。
   并且**每一个非「已解析」的状态都必须带原因**（纪律第 9 条的另一面：
   判出来了却不说为什么，等于没判）。

3. 第二次扫描要能认出新增/修改/消失，且消失的文件只标记、不删行。

测试数据一律建在 pytest 的 tmp_path 下，并用 FJHY-TEST-DOCLIB- 前缀登记，
跑完删干净（纪律第 12 条：验证动作产生的数据要一眼可辨、事后好清理）。
"""
import datetime
import hashlib
import os
import sys
import time
import zipfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import db, doclib  # noqa: E402

PREFIX = "FJHY-TEST-DOCLIB-"


# ─────────────────────────────────────────────────────────────────────────────
# 工具
# ─────────────────────────────────────────────────────────────────────────────
def snapshot_dir(root):
    """目录的完整指纹：路径 + 大小 + 修改时间 + 内容摘要。用来验「只读」。"""
    out = {}
    for dirpath, dirnames, filenames in os.walk(root):
        for fn in filenames:
            p = os.path.join(dirpath, fn)
            st = os.stat(p)
            with open(p, "rb") as f:
                digest = hashlib.sha1(f.read()).hexdigest()
            out[os.path.relpath(p, root)] = (st.st_size, int(st.st_mtime), digest)
    return out


def make_pdf(path, text=None):
    """造 PDF：text 为空就是一页空白（没有文字层），模拟纯扫描件。"""
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 120), text, fontsize=12, fontname="china-s")
    doc.save(path)
    doc.close()


def make_docx(path, paras, table=None):
    import docx
    d = docx.Document()
    for t in paras:
        d.add_paragraph(t)
    if table:
        tb = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, v in enumerate(row):
                tb.cell(i, j).text = str(v)
    d.save(path)


@pytest.fixture
def lib(tmp_path):
    """一个登记好的临时文件库根。跑完连带删掉登记（磁盘目录 pytest 自己收）。"""
    root = tmp_path / "资料库"
    (root / "制度").mkdir(parents=True)
    (root / "制度" / "差旅管理办法.txt").write_text(
        "第一条 为规范差旅费管理，制定本办法。\n第二条 住宿费标准见附件。",
        encoding="utf-8")
    res = doclib.add_roots([str(root)], label=PREFIX + "资料库")
    assert res[0]["ok"], res
    rid = res[0]["root_id"]
    yield root, rid
    db.execute("DELETE FROM doc_lib_change WHERE abs_path LIKE %s", (str(root) + "%",))
    db.execute("DELETE FROM doc_lib_root WHERE root_id=%s", (rid,))


# ─────────────────────────────────────────────────────────────────────────────
# 纪律一 · 只读
# ─────────────────────────────────────────────────────────────────────────────
def test_扫描不写入被扫目录(lib):
    root, rid = lib
    make_pdf(str(root / "制度" / "通知.pdf"), "关于印发差旅费管理办法的通知")
    make_docx(str(root / "制度" / "标准表.docx"), ["住宿费标准"],
              [["城市", "处级", "科级"], ["福州", "460", "420"]])
    before = snapshot_dir(str(root))
    doclib.scan([rid])
    after = snapshot_dir(str(root))
    assert before == after, "扫描动了被扫目录：%s" % (
        set(before.items()) ^ set(after.items()))


def test_移出文件库不碰磁盘(lib):
    root, rid = lib
    doclib.scan([rid])
    before = snapshot_dir(str(root))
    r = doclib.remove_root(rid)
    assert r["removed"]
    assert snapshot_dir(str(root)) == before
    # 登记没了，文件还在
    assert (root / "制度" / "差旅管理办法.txt").exists()
    assert db.scalar("SELECT count(*) FROM doc_lib_file WHERE root_id=%s", (rid,)) == 0


# ─────────────────────────────────────────────────────────────────────────────
# 纪律二 · 六态不许合并，且必须带原因
# ─────────────────────────────────────────────────────────────────────────────
def test_纯扫描PDF判需OCR而不是未解析(tmp_path):
    p = tmp_path / "扫描件.pdf"
    make_pdf(str(p))                       # 一页空白 = 没有文字层
    state, body, reason = doclib.extract_text(str(p))
    assert state == "需OCR", "纯扫描件被判成了 %s：%s" % (state, reason)
    assert body == ""
    assert "OCR" in reason and "文字层" in reason


def test_有文字层PDF直接取文(tmp_path):
    p = tmp_path / "通知.pdf"
    make_pdf(str(p), "关于印发差旅费管理办法的通知")
    state, body, reason = doclib.extract_text(str(p))
    assert state == "已解析", reason
    assert "差旅费" in body


def test_docx必须取到表格内容(tmp_path):
    """住宿费标准、交通工具等级这类最关键的东西都在表里。
    只取段落会解析出一堆空话，状态却显示「已解析」——比解析失败更难发现。"""
    p = tmp_path / "标准.docx"
    make_docx(str(p), ["附件1 住宿费标准"],
              [["城市", "处级", "科级"], ["福州", "460", "420"]])
    state, body, _ = doclib.extract_text(str(p))
    assert state == "已解析"
    assert "住宿费标准" in body
    assert "460" in body and "福州" in body, "表格内容没取到：%r" % body


def test_老格式判未解析且原因要写下一步(tmp_path):
    for ext in (".doc", ".xls", ".ofd", ".wps", ".et", ".rtf"):
        p = tmp_path / ("制度" + ext)
        p.write_bytes(b"\xd0\xcf\x11\xe0" + b"x" * 200)
        state, body, reason = doclib.extract_text(str(p))
        assert state == "未解析", "%s 判成了 %s" % (ext, state)
        assert body == ""
        assert reason and ("另存为" in reason or "导出" in reason), \
            "%s 的原因没写下一步该干什么：%s" % (ext, reason)


def test_空文件与解析失败是两件事(tmp_path):
    empty = tmp_path / "空.txt"
    empty.write_bytes(b"")
    assert doclib.extract_text(str(empty))[0] == "空文件"

    fake = tmp_path / "其实不是docx.docx"      # 改过扩展名的垃圾
    fake.write_bytes(b"not a zip at all")
    state, _, reason = doclib.extract_text(str(fake))
    assert state == "解析失败", state
    assert reason


def test_每个非已解析状态都带原因(tmp_path):
    """纪律第 9 条的另一面：判出来了却不说为什么，等于没判。"""
    cases = []
    e = tmp_path / "e.txt"; e.write_bytes(b""); cases.append(e)
    d = tmp_path / "d.doc"; d.write_bytes(b"xx"); cases.append(d)
    s = tmp_path / "s.pdf"; make_pdf(str(s)); cases.append(s)
    f = tmp_path / "f.docx"; f.write_bytes(b"junk"); cases.append(f)
    for p in cases:
        state, _, reason = doclib.extract_text(str(p))
        assert state != "已解析"
        assert reason and reason.strip(), "%s 判成 %s 却没给原因" % (p.name, state)


def test_编码是gb18030的中文txt能读(tmp_path):
    p = tmp_path / "gbk.txt"
    p.write_bytes("第十条 市内交通费每人每天80元。".encode("gb18030"))
    state, body, reason = doclib.extract_text(str(p))
    assert state == "已解析"
    assert "市内交通费" in body, body


# ─────────────────────────────────────────────────────────────────────────────
# 纪律三 · 变化感知
# ─────────────────────────────────────────────────────────────────────────────
def test_第二次扫描认出新增修改消失(lib):
    root, rid = lib
    first = doclib.scan([rid])
    assert first["added"] == 1 and first["updated"] == 0
    assert first["missing"] == 0

    # 第二次：什么都没动 → 全部未变，台账里一条变化都不该有
    again = doclib.scan([rid])
    assert again["added"] == 0 and again["updated"] == 0 and again["missing"] == 0
    assert again["unchanged"] == 1
    assert doclib.list_changes(again["scan_id"])["total"] == 0, \
        "文件没动却记了变化，台账会被假变更污染"

    # 第三次：加一份、改一份
    (root / "制度" / "报销细则.txt").write_text("第一条 报销需附行程单。", encoding="utf-8")
    time.sleep(1.1)                                  # 让 mtime 确实变化（秒级精度）
    (root / "制度" / "差旅管理办法.txt").write_text(
        "第一条 为规范差旅费管理，制定本办法。\n第二条 住宿费标准见附件。\n第三条 新增条款。",
        encoding="utf-8")
    third = doclib.scan([rid])
    assert third["added"] == 1, third
    assert third["updated"] == 1, third
    kinds = {c["change"] for c in third["changes"]}
    assert kinds == {"新增", "修改"}, kinds
    detail = [c["detail"] for c in third["changes"] if c["change"] == "修改"][0]
    assert "大小" in detail or "修改时间" in detail, detail

    # 第四次：删一份 → 标消失，但登记行必须还在
    os.remove(root / "制度" / "报销细则.txt")
    fourth = doclib.scan([rid])
    assert fourth["missing"] == 1, fourth
    row = db.query_one("SELECT missing, last_change FROM doc_lib_file WHERE file_name=%s",
                       ("报销细则.txt",))
    assert row is not None, "消失的文件被删了行——就再也说不清规则当初依据的是哪份文件"
    assert row["missing"] is True and row["last_change"] == "消失"


def test_强制重扫不算修改(lib):
    """文件没动，强制重扫只是重跑解析器。记成「修改」会把台账污染成一片假变更。"""
    root, rid = lib
    doclib.scan([rid])
    forced = doclib.scan([rid], force=True)
    assert forced["updated"] == 0, forced
    assert doclib.list_changes(forced["scan_id"])["total"] == 0
    assert forced["mode"] == "强制重扫"


def test_只在扫过的根里判消失(lib, tmp_path):
    """没扫的根里的文件当然看不见，把它们标成消失是纯粹的误报。"""
    root, rid = lib
    doclib.scan([rid])
    other = tmp_path / "另一个库"
    other.mkdir()
    (other / "别的制度.txt").write_text("内容", encoding="utf-8")
    r2 = doclib.add_roots([str(other)], label=PREFIX + "另一个库")[0]
    try:
        doclib.scan([r2["root_id"]])          # 只扫第二个根
        n = db.scalar("SELECT count(*) FROM doc_lib_file WHERE root_id=%s AND missing",
                      (rid,))
        assert n == 0, "扫 B 根却把 A 根的文件标成了消失"
    finally:
        db.execute("DELETE FROM doc_lib_change WHERE abs_path LIKE %s", (str(other) + "%",))
        db.execute("DELETE FROM doc_lib_root WHERE root_id=%s", (r2["root_id"],))


# ─────────────────────────────────────────────────────────────────────────────
# 待抽取：用指纹比对算，不设标志位
# ─────────────────────────────────────────────────────────────────────────────
def test_待抽取随文件变化自动进出(lib):
    root, rid = lib
    doclib.scan([rid])
    fid = db.scalar("SELECT file_id FROM doc_lib_file WHERE root_id=%s", (rid,))

    ids = [r["file_id"] for r in doclib.pending_extract()["items"]]
    assert fid in ids, "新扫进来的文件不在待抽取里"

    doclib.mark_extracted([fid])
    ids = [r["file_id"] for r in doclib.pending_extract()["items"]]
    assert fid not in ids, "抽过了还在待抽取里"

    time.sleep(1.1)
    (root / "制度" / "差旅管理办法.txt").write_text("改了内容", encoding="utf-8")
    doclib.scan([rid])
    rows = {r["file_id"]: r for r in doclib.pending_extract()["items"]}
    assert fid in rows, "抽取后文件又被改过，却没回到待抽取——规则会悄悄对不上现行制度"
    assert rows[fid]["why"] == "抽取后文件又被改过"


def test_取不到正文的文件不算待抽取但要报数(lib):
    """「抽完了」不等于「制度都抽过了」。取不到正文的那批必须单独报出来。"""
    root, rid = lib
    make_pdf(str(root / "制度" / "扫描件.pdf"))      # 没有文字层
    doclib.scan([rid])
    p = doclib.pending_extract()
    assert all(r["parse_state"] == "已解析" for r in p["items"])
    states = {b["parse_state"]: b["n"] for b in p["blocked"]}
    assert states.get("需OCR", 0) >= 1, states


# ─────────────────────────────────────────────────────────────────────────────
# 路径安全与闸门
# ─────────────────────────────────────────────────────────────────────────────
def test_路径非法一律抛PathError(tmp_path):
    f = tmp_path / "不是文件夹.txt"
    f.write_text("x", encoding="utf-8")
    for bad in ("", "   ", str(tmp_path / "不存在的目录"), str(f)):
        with pytest.raises(doclib.PathError):
            doclib.safe_dir(bad)


def test_office锁文件不进库(lib):
    """~$xxx.docx 扩展名合法、内容是几百字节二进制，收进来只会多一批解析失败。"""
    root, rid = lib
    (root / "制度" / "~$差旅管理办法.docx").write_bytes(b"\x00" * 160)
    doclib.scan([rid])
    n = db.scalar("SELECT count(*) FROM doc_lib_file WHERE file_name LIKE %s", ("~$%",))
    assert n == 0


def test_非文档扩展名不进库(lib):
    root, rid = lib
    (root / "制度" / "工具.exe").write_bytes(b"MZ" + b"\x00" * 100)
    (root / "制度" / "照片.jpg").write_bytes(b"\xff\xd8\xff" + b"\x00" * 100)
    doclib.scan([rid])
    exts = {r["ext"] for r in doclib.list_files(root_id=rid)}
    assert ".exe" not in exts and ".jpg" not in exts, exts


# ─────────────────────────────────────────────────────────────────────────────
# 文件级属性：分类 / 文号 / 标题
# ─────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("name,kind", [
    # 制度
    ("2.1关于印发《差旅管理实施细则》的通知（闽港财智〔2025〕20号）.pdf", "制度"),
    ("2.1国资委差旅制度.pdf", "制度"),
    ("2.3 差旅费共享审核要点.pdf", "制度"),
    ("差旅费审核规则.xlsx", "制度"),
    ("2.1差旅费解答六（国资委未转发）.pdf", "制度"),
    # 单据：实扫 14 份客户资料里有 5 份是这个。
    # 混进「制度」统计，「还有多少条没变成规则」这个数字就失真了。
    ("多人差旅报账单明细表.xlsx", "单据"),
    ("差旅费审批意见1.pdf", "单据"),
    ("差旅报账单明细表.xlsx", "单据"),
    # 范本：文书模板，里面的「第X条」是文书条款不是制度条款
    ("借款申请书模板.docx", "范本"),
    ("特别约定.docx", "范本"),
    # 判不出就是「其他」，不猜
    ("2.5 系统审核算法.xlsx", "其他"),
    ("扫描件001.pdf", "其他"),
])
def test_文件分类(name, kind):
    assert doclib.classify(name) == kind, doclib.classify(name)


def test_范本优先于单据():
    """「出差审批表单」是空白模板按范本算，「差旅费审批意见」是具体某一单按单据算。
    顺序反过来，模板就会被当成真单据。"""
    assert doclib.classify("出差审批单模板.xlsx") == "范本"
    assert doclib.classify("差旅费审批意见.pdf") == "单据"


@pytest.mark.parametrize("name,no,year", [
    ("2.1关于印发《差旅管理实施细则》的通知（闽港财智〔2025〕20号）.pdf",
     "闽港财智〔2025〕20号", 2025),
    ("2.1[2024003]闽港集团纪〔2023〕23号关于做好监督检查整改工作的通知.pdf",
     "闽港集团纪〔2023〕23号", 2023),
    ("关于印发某办法的通知(闽财行[2017]17号).pdf", "闽财行[2017]17号", 2017),
])
def test_从文件名取文号(name, no, year):
    got_no, _, got_year = doclib.parse_doc_meta(name)
    assert got_no == no, got_no
    assert got_year == year


def test_没有文号就留空不猜():
    """**取不到就空着。** 项目纪律：溯源只认硬证据。
    猜一个文号出来，下游会拿它去匹配规则，匹配上的是另一份文件——
    错配给出的是一条看起来有据可查、实际指错地方的证据链，
    比标「未匹配」危险得多。"""
    for n in ("2.1国资委差旅制度.pdf", "差旅费审核规则.xlsx", "扫描件001.pdf",
              "会议纪要2025.docx", "报销标准20250101.pdf"):
        no, title, year = doclib.parse_doc_meta(n)
        assert no is None, "%s 猜出了文号 %s" % (n, no)
        assert year is None


def test_多个书名号取最长的那个():
    """「关于印发《业务接待管理办法》的通知」里，要的是被印发的那份办法，
    不是「通知」本身。"""
    _, t, _ = doclib.parse_doc_meta(
        "2.1关于印发《业务接待管理办法》的通知（闽港财智〔2025〕18号）.pdf")
    assert t == "业务接待管理办法", t
    _, t2, _ = doclib.parse_doc_meta("《甲》与《乙类管理办法》的衔接.docx")
    assert t2 == "乙类管理办法", t2


def test_扫描后文件级属性入库(lib):
    root, rid = lib
    (root / "制度" / "关于印发《测试管理办法》的通知（闽港财智〔2025〕99号）.txt").write_text(
        "第一条 测试。", encoding="utf-8")
    doclib.scan([rid])
    row = db.query_one("SELECT doc_kind, doc_number, doc_title, issued_year "
                       "FROM doc_lib_file WHERE file_name LIKE %s", ("%测试管理办法%",))
    assert row["doc_kind"] == "制度"
    assert row["doc_number"] == "闽港财智〔2025〕99号"
    assert row["doc_title"] == "测试管理办法"
    assert row["issued_year"] == 2025
